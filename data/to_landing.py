#!/usr/bin/env python3
"""Convert the generator output in ``data/out`` into landing CSVs under
``data/out/landing``, ready to upload to the Snowflake ``RAW`` schema.

This mirrors how Azure Data Factory lands source files at work: everything
arrives in RAW *exactly as it is*, with no cleaning. The messy spreadsheet is
flattened to CSV without trimming, parsing, or deduping anything — those are
the next task's job.

* ``circuits.csv``, ``districts.csv``, ``vendors.csv``, ``answer_key.csv``
  are copied through byte-for-byte.
* ``row_clearing_plan.xlsx`` becomes ``row_clearing_plan.csv``:
  - a ``source_row`` column with the Excel row number (matches ``excel_row``
    in the answer key)
  - every cell as raw text, exactly as it appears (no trimming, no type
    conversion); blank cells become an empty string
  - real Excel date cells are written as ISO ``YYYY-MM-DD``
  - each of the two date columns gets a cell-type column
    (``date``, ``text``, or ``blank``) so a typed ``"2027-03-15"`` can be told
    apart from a real date cell

Deterministic: the same input produces identical landing files on every run.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import shutil
from pathlib import Path

from openpyxl import load_workbook

# Files copied through unchanged, in a stable order.
PASSTHROUGH = ["circuits.csv", "districts.csv", "vendors.csv", "answer_key.csv"]

# Header of the landing CSV built from the spreadsheet. The two date columns
# each carry an adjacent cell-type column.
PLAN_HEADER = [
    "source_row",
    "feeder_id",
    "last_cleared",
    "last_cleared_celltype",
    "last_work_done",
    "next_clearing",
    "next_clearing_celltype",
    "vendor",
    "notes",
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--in", dest="in_dir", type=Path, default=Path("data/out"),
                   help="input directory (default data/out)")
    p.add_argument("--out", dest="out_dir", type=Path,
                   default=Path("data/out/landing"),
                   help="output directory (default data/out/landing)")
    return p.parse_args()


def cell_text_and_type(value) -> tuple[str, str]:
    """Return (raw text, cell-type) for one spreadsheet cell.

    - blank cell            -> ("", "blank")
    - real date/datetime    -> ("YYYY-MM-DD", "date")
    - anything else (text)  -> (the value verbatim, "text")

    Nothing is trimmed: leading and trailing spaces are preserved exactly.
    """
    if value is None:
        return "", "blank"
    if isinstance(value, dt.datetime):
        return value.date().isoformat(), "date"
    if isinstance(value, dt.date):
        return value.isoformat(), "date"
    return str(value), "text"


def plan_text(value) -> str:
    """Raw text for a non-date cell: verbatim, blanks as empty string."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def convert_plan(xlsx_path: Path, csv_path: Path) -> int:
    """Flatten the clearing-plan spreadsheet to a landing CSV. Returns the
    number of data rows written."""
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    # Column positions in the spreadsheet (1-based):
    # 1 Feeder ID | 2 Last Cleared | 3 Last Work Done | 4 Next Clearing
    # 5 Vendor | 6 Notes
    n = 0
    # QUOTE_ALL keeps every value exact, including leading/trailing spaces and
    # empty strings; newline="" + \n makes the output byte-stable across OSes.
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, quoting=csv.QUOTE_ALL, lineterminator="\n")
        w.writerow(PLAN_HEADER)
        for r in range(2, ws.max_row + 1):
            feeder = plan_text(ws.cell(row=r, column=1).value)
            last_cleared, last_type = cell_text_and_type(ws.cell(row=r, column=2).value)
            last_work = plan_text(ws.cell(row=r, column=3).value)
            next_clearing, next_type = cell_text_and_type(ws.cell(row=r, column=4).value)
            vendor = plan_text(ws.cell(row=r, column=5).value)
            notes = plan_text(ws.cell(row=r, column=6).value)
            w.writerow([
                r, feeder,
                last_cleared, last_type,
                last_work,
                next_clearing, next_type,
                vendor, notes,
            ])
            n += 1
    return n


def main() -> None:
    args = parse_args()
    in_dir: Path = args.in_dir
    out_dir: Path = args.out_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"in: {in_dir}   out: {out_dir}")

    for name in PASSTHROUGH:
        src = in_dir / name
        shutil.copyfile(src, out_dir / name)
        print(f"  copied   {name}")

    n = convert_plan(in_dir / "row_clearing_plan.xlsx",
                     out_dir / "row_clearing_plan.csv")
    print(f"  wrote    row_clearing_plan.csv  {n} data rows")


if __name__ == "__main__":
    main()
